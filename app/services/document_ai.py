"""
Google Document AI extraction service
"""
from google.cloud import documentai_v1 as documentai
from google.api_core.client_options import ClientOptions
from typing import List, Dict, Any, Optional
from app.config import get_settings
from app.utils.logger import get_logger
from app.services.cost_tracker import CostTracker
import io
import os
from pypdf import PdfReader, PdfWriter
import docx

settings = get_settings()
logger = get_logger(__name__)


class DocumentAIExtractor:
    """
    Service for extracting structured content from documents using GCP Document AI
    """
    
    def __init__(self):
        self.project_id = settings.gcp_project_id
        self.location = settings.docai_location
        self.layout_processor_id = settings.docai_layout_processor_id
        self.form_processor_id = settings.docai_form_processor_id
        
        # Initialize Document AI client
        opts = ClientOptions(api_endpoint=f"{self.location}-documentai.googleapis.com")
        self.client = documentai.DocumentProcessorServiceClient(client_options=opts)
    
    async def extract_document(
        self,
        file_path: str,
        mime_type: str = "application/pdf",
        cost_tracker: CostTracker = None,
    ) -> Dict[str, Any]:
        """
        Extract document structure using Document AI with large file support
        """
        logger.info(f"Starting document extraction: {file_path}")
        
        # Reject legacy .doc files - not supported by Document AI
        if mime_type == "application/msword":
            logger.warning(f"Legacy .doc file rejected: {file_path}")
            raise ValueError(
                "Legacy Microsoft Word .doc format is not supported. "
                "Please convert your file to .docx format and upload again. "
                "You can convert using Microsoft Word (Save As > .docx) or "
                "online converters like CloudConvert."
            )
        
        # Read file
        with open(file_path, "rb") as f:
            file_content = f.read()
        
        # Validate file content (especially for PDFs)
        if len(file_content) == 0:
            raise ValueError(
                f"Uploaded file is empty (0 bytes). Please check the file and try again."
            )
        
        logger.info(f"File size: {len(file_content)} bytes")
            
        # Check if we need to split (PDF only)
        if mime_type == "application/pdf":
            try:
                reader = PdfReader(io.BytesIO(file_content))
                num_pages = len(reader.pages)
                
                logger.info(f"PDF validation: {num_pages} pages detected")
                
                if num_pages == 0:
                    raise ValueError(
                        "PDF file appears to be empty or corrupted. "
                        "Please verify the file can be opened in a PDF reader."
                    )
                
                # Check if PDF has text or is just images (scanned)
                first_page_text = reader.pages[0].extract_text().strip()
                if not first_page_text and num_pages > 0:
                    logger.warning(f"PDF appears to be scanned images (no text layer on page 1). Document AI will attempt OCR.")
                
                if num_pages > 30:  # Document AI Online limit is usually 15-30
                    logger.info(f"Large document detected ({num_pages} pages). Splitting into chunks...")
                    large_result = await self._process_large_pdf(reader, mime_type)
                    # If we got content, return it.
                    if large_result.get('elements') or large_result.get('full_text'):
                        return large_result
                    
                    # If local/chunked processing failed, jump directly to PyPDF fallback 
                    # for the ENTIRE document using the reader we already have.
                    logger.warning("Large PDF processing returned no content from Document AI. Proceeding to PyPDF fallback.")
                    fallback_elements, fallback_text = await self._extract_pdf_fallback(file_path, reader=reader)
                    
                    return {
                        "total_pages": num_pages,
                        "elements": fallback_elements,
                        "full_text": fallback_text,
                        "raw_layout": None,
                        "raw_form": None,
                        "is_merged": True
                    }

            except Exception as e:
                logger.error(f"Error reading PDF: {e}")
                raise ValueError(
                    f"Unable to read PDF file. The file may be corrupted, password-protected, or in an unsupported format. Error: {str(e)}"
                )
        
        # Reader might be available from above (for large PDF check)
        # But we'll still need content for the parsers below
        layout_result = await self._process_with_layout_parser(file_content, mime_type)
        
        # Track layout parser page cost
        if cost_tracker and hasattr(layout_result, 'pages'):
            layout_pages = len(layout_result.pages)
            cost_tracker.record_docai_pages(layout_pages, parser_type="layout")
        
        # Process with form parser for enhanced field recognition (Optional enhancement)
        form_result = None
        try:
            form_result = await self._process_with_form_parser(file_content, mime_type)
            # Track form parser page cost
            if cost_tracker and form_result and hasattr(form_result, 'pages'):
                form_pages = len(form_result.pages)
                cost_tracker.record_docai_pages(form_pages, parser_type="form")
        except Exception as e:
            logger.warning(f"Optional form parser failed for {mime_type}: {e}. Proceeding with layout parser only.")
        
        # Combine and structure results
        elements = self._build_structured_elements(layout_result, form_result)
        full_text = layout_result.text if hasattr(layout_result, 'text') else ""
        
        # Fallback: If Document AI returns nothing, try native extraction
        if len(elements) == 0 and len(full_text) == 0:
            logger.warning("Document AI returned no content. Attempting fallback extraction...")
            
            if mime_type == "application/pdf":
                logger.info("Using PyPDF fallback for PDF extraction")
                # reader might be locally scoped in extract_document, but we'll try to get it again or re-read
                fallback_elements, fallback_text = await self._extract_pdf_fallback(file_path)
                if fallback_elements or fallback_text:
                    elements = fallback_elements
                    full_text = fallback_text
                    logger.info(f"PyPDF fallback successful: {len(elements)} elements, {len(fallback_text)} chars")
                else:
                    logger.error("PyPDF fallback also failed")
            
            elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or file_path.endswith('.docx'):
                logger.info("Using native DOCX fallback")
                native_elements, native_text = self._extract_docx_native(file_path)
                if native_elements:
                    elements = native_elements
                    full_text = native_text or full_text
                    logger.info(f"DOCX fallback successful: {len(elements)} elements")
        
        logger.info(f"Extraction complete: {len(elements)} elements extracted, {len(full_text)} characters of text found")
        
        # Better error reporting when nothing is extracted after all fallbacks
        if len(elements) == 0 and len(full_text) == 0:
            error_msg = (
                "Could not extract any content from this file using Document AI or fallback methods. "
                "Possible causes:\n"
                "- Document AI processor may be misconfigured (check processor ID in .env)\n"
                "- PDF is a scanned image without text (OCR failed)\n"
                "- File is corrupted or encrypted\n"
                "Please try:\n"
                "1. Verify Document AI processor is set up correctly in GCP Console\n"
                "2. If scanned PDF, re-scan with higher quality or use OCR software first\n"
                "3. Convert to .docx format if possible"
            )
            logger.error(error_msg)
            raise ValueError(error_msg)
        
        if len(elements) == 0 and len(full_text) > 0:
            logger.warning("No elements extracted but text was found. This may happen with some DOCX files or processors.")
        
        return {
            "total_pages": len(layout_result.pages) if hasattr(layout_result, 'pages') else 1,
            "elements": elements,
            "full_text": full_text,
            "raw_layout": layout_result,
            "raw_form": form_result
        }
    
    async def _process_large_pdf(
        self,
        reader: PdfReader,
        mime_type: str,
        chunk_size: int = 15
    ) -> Dict[str, Any]:
        """Process large PDF by splitting into chunks and merging results"""
        
        num_pages = len(reader.pages)
        layout_docs = []
        form_docs = []
        all_elements = []
        combined_text = ""
        current_offset = 0
        
        for i in range(0, num_pages, chunk_size):
            end_page = min(i + chunk_size, num_pages)
            logger.info(f"Processing PDF chunk: pages {i+1} to {end_page}")
            
            # Create chunk
            writer = PdfWriter()
            for page_num in range(i, end_page):
                writer.add_page(reader.pages[page_num])
            
            chunk_buffer = io.BytesIO()
            writer.write(chunk_buffer)
            chunk_content = chunk_buffer.getvalue()
            
            # Process chunk (Layout is required, Form is optional)
            layout_doc = await self._process_with_layout_parser(chunk_content, mime_type)
            
            form_doc = None
            try:
                form_doc = await self._process_with_form_parser(chunk_content, mime_type)
            except Exception as e:
                logger.warning(f"Optional form parser failed for chunk: {e}")
            
            # Extract elements from this chunk
            chunk_elements = self._build_structured_elements(
                layout_doc, form_doc, page_offset=i
            )
            
            # Extraction complete for this chunk
            layout_docs.append(layout_doc)
            form_docs.append(form_doc)
            all_elements.extend(chunk_elements)
            combined_text += layout_doc.text if hasattr(layout_doc, 'text') else ""
            
        logger.info(f"Successfully processed all {len(layout_docs)} chunks")
        
        return {
            "total_pages": num_pages,
            "elements": all_elements,
            "full_text": combined_text,
            "raw_layout": layout_docs[0],  # Return first chunk as representative
            "raw_form": form_docs[0],
            "is_merged": True
        }
    
    async def _process_with_layout_parser(
        self,
        file_content: bytes,
        mime_type: str
    ) -> documentai.Document:
        """Process document with layout parser"""
        logger.info(f"Calling Layout Parser for mime_type: {mime_type}")
        
        processor_name = self.client.processor_path(
            self.project_id,
            self.location,
            self.layout_processor_id
        )
        
        logger.info(f"Processor path: {processor_name}")
        logger.info(f"Sending {len(file_content)} bytes to Document AI")
        
        raw_document = documentai.RawDocument(content=file_content, mime_type=mime_type)
        request = documentai.ProcessRequest(name=processor_name, raw_document=raw_document)
        
        result = self.client.process_document(request=request)
        doc = result.document
        
        # Detailed logging of response
        logger.info(f"Layout Parser response: {len(doc.pages)} pages, {len(doc.text)} characters")
        
        if len(doc.pages) == 0:
            logger.error(f"Document AI returned 0 pages!")
            logger.error(f"API Response details:")
            logger.error(f"  - MIME type sent: {mime_type}")
            logger.error(f"  - Bytes sent: {len(file_content)}")
            logger.error(f"  - Processor: {processor_name}")
            logger.error(f"  - Has document object: {doc is not None}")
            logger.error(f"  - Document text length: {len(doc.text) if hasattr(doc, 'text') else 'N/A'}")
            logger.error(f"  - Document entities: {len(doc.entities) if hasattr(doc, 'entities') else 'N/A'}")
            
            # Check if there's error info in the response
            if hasattr(result, 'error'):
                logger.error(f"  - API Error: {result.error}")
            
            # Try Form Parser as OCR fallback before PyPDF
            if self.form_processor_id:
                try:
                    logger.info("Layout Parser returned 0 pages. Attempting Form Parser as OCR fallback...")
                    form_doc = await self._process_with_form_parser(file_content, mime_type)
                    if len(form_doc.pages) > 0:
                        logger.info(f"Form Parser OCR fallback succeeded: {len(form_doc.pages)} pages, {len(form_doc.text)} characters")
                        return form_doc
                    else:
                        logger.warning("Form Parser also returned 0 pages")
                except Exception as e:
                    logger.warning(f"Form Parser OCR fallback failed: {e}")
        
        return doc
    
    async def _process_with_form_parser(
        self,
        file_content: bytes,
        mime_type: str
    ) -> documentai.Document:
        """Process document with form parser for enhanced field extraction"""
        
        processor_name = self.client.processor_path(
            self.project_id,
            self.location,
            self.form_processor_id
        )
        
        raw_document = documentai.RawDocument(content=file_content, mime_type=mime_type)
        request = documentai.ProcessRequest(name=processor_name, raw_document=raw_document)
        
        result = self.client.process_document(request=request)
        return result.document
    
    def _build_structured_elements(
        self,
        layout_doc: documentai.Document,
        form_doc: Optional[documentai.Document] = None,
        page_offset: int = 0
    ) -> List[Dict[str, Any]]:
        """
        Build structured elements from Document AI output
        """
        elements = []
        order_index = 0
        
        # Use pages from layout_doc as primary source
        logger.info(f"Building elements from {len(layout_doc.pages)} pages")
        
        for page_idx, page in enumerate(layout_doc.pages):
            actual_page_num = page_idx + 1 + page_offset
            
            logger.debug(f"Processing page {actual_page_num}: {len(page.tables)} tables, {len(page.paragraphs)} paragraphs, {len(page.blocks)} blocks")
            
            # If form_doc is available, we could merge fields, 
            # but primary fields usually exist in layout_doc too in newer versions.
            # For simplicity, we use layout_doc as source.
            
            # Extract tables
            for table in page.tables:
                element = self._extract_table_element(
                    table, actual_page_num, order_index, layout_doc.text
                )
                elements.append(element)
                order_index += 1
            
            # Extract form fields
            for form_field in page.form_fields:
                element = self._extract_form_element(
                    form_field, actual_page_num, order_index, layout_doc.text
                )
                elements.append(element)
                order_index += 1
            
            # Extract paragraphs and text blocks
            # Fallback to page.blocks if page.paragraphs is empty
            text_containers = page.paragraphs if page.paragraphs else page.blocks
            
            for container in text_containers:
                element = self._extract_text_element(
                    container, actual_page_num, order_index, layout_doc.text
                )
                elements.append(element)
                order_index += 1
            
            # Extract list items
            for line in page.lines:
                if self._is_list_item(line, layout_doc.text):
                    element = self._extract_list_element(
                        line, actual_page_num, order_index, layout_doc.text
                    )
                    elements.append(element)
                    order_index += 1
        
        # Build hierarchy
        elements = self._build_hierarchy(elements)
        
        return elements
    
    def _extract_table_element(
        self,
        table: documentai.Document.Page.Table,
        page_num: int,
        order_index: int,
        full_text: str
    ) -> Dict[str, Any]:
        """Extract table element with structure preservation"""
        
        # Build table data
        rows = []
        for row in table.body_rows:
            cells = []
            for cell in row.cells:
                cell_text = self._get_text_from_layout(cell.layout, full_text)
                cells.append({
                    "text": cell_text,
                    "row_span": cell.row_span,
                    "col_span": cell.col_span
                })
            rows.append(cells)
        
        # Extract headers
        headers = []
        for header_row in table.header_rows:
            header_cells = []
            for cell in header_row.cells:
                cell_text = self._get_text_from_layout(cell.layout, full_text)
                header_cells.append(cell_text)
            headers.append(header_cells)
        
        return {
            "element_id": f"table_{page_num}_{order_index}",
            "element_type": "table",
            "page_number": page_num,
            "order_index": order_index,
            "bounding_box": self._extract_bounding_box(table.layout),
            "content": {
                "headers": headers,
                "rows": rows,
                "row_count": len(rows),
                "col_count": len(rows[0]) if rows else 0
            }
        }
    
    def _extract_form_element(
        self,
        form_field: documentai.Document.Page.FormField,
        page_num: int,
        order_index: int,
        full_text: str
    ) -> Dict[str, Any]:
        """Extract form field element"""
        
        field_name = self._get_text_from_layout(form_field.field_name, full_text)
        field_value = self._get_text_from_layout(form_field.field_value, full_text)
        
        return {
            "element_id": f"form_{page_num}_{order_index}",
            "element_type": "form",
            "page_number": page_num,
            "order_index": order_index,
            "bounding_box": self._extract_bounding_box(form_field.field_name),
            "content": {
                "field_name": field_name,
                "field_value": field_value,
                "confidence": form_field.field_name.confidence if hasattr(form_field.field_name, 'confidence') else None
            }
        }
    
    def _extract_text_element(
        self,
        paragraph: Any,  # Can be Paragraph or Block
        page_num: int,
        order_index: int,
        full_text: str
    ) -> Dict[str, Any]:
        """Extract text paragraph element"""
        
        text = self._get_text_from_layout(paragraph.layout, full_text)
        
        return {
            "element_id": f"text_{page_num}_{order_index}",
            "element_type": "text",
            "page_number": page_num,
            "order_index": order_index,
            "bounding_box": self._extract_bounding_box(paragraph.layout),
            "content": {
                "text": text
            }
        }
    
    def _extract_list_element(
        self,
        line: documentai.Document.Page.Line,
        page_num: int,
        order_index: int,
        full_text: str
    ) -> Dict[str, Any]:
        """Extract list item element"""
        
        text = self._get_text_from_layout(line.layout, full_text)
        
        return {
            "element_id": f"list_{page_num}_{order_index}",
            "element_type": "list_item",
            "page_number": page_num,
            "order_index": order_index,
            "bounding_box": self._extract_bounding_box(line.layout),
            "content": {
                "text": text
            }
        }
    
    def _get_text_from_layout(
        self,
        layout: documentai.Document.Page.Layout,
        full_text: str
    ) -> str:
        """Extract text from layout using text anchors"""
        
        if not layout.text_anchor.text_segments:
            return ""
        
        text_parts = []
        for segment in layout.text_anchor.text_segments:
            start = int(segment.start_index) if segment.start_index else 0
            end = int(segment.end_index) if segment.end_index else len(full_text)
            text_parts.append(full_text[start:end])
        
        return "".join(text_parts).strip()
    
    def _extract_bounding_box(
        self,
        layout: documentai.Document.Page.Layout
    ) -> Dict[str, Any]:
        """Extract bounding box coordinates"""
        
        if not layout.bounding_poly:
            return {}
        
        vertices = []
        normalized_vertices = []
        
        for vertex in layout.bounding_poly.vertices:
            vertices.append([vertex.x, vertex.y])
        
        for vertex in layout.bounding_poly.normalized_vertices:
            normalized_vertices.append([vertex.x, vertex.y])
        
        return {
            "vertices": vertices,
            "normalized_vertices": normalized_vertices
        }
    
    def _is_list_item(
        self,
        line: documentai.Document.Page.Line,
        full_text: str
    ) -> bool:
        """Determine if a line is a list item"""
        
        text = self._get_text_from_layout(line.layout, full_text)
        
        # Simple heuristic: starts with bullet, number, or letter
        list_markers = ["•", "◦", "▪", "-", "–", "—"]
        if any(text.startswith(marker) for marker in list_markers):
            return True
        
        # Check for numbered lists (1., a., i., etc.)
        import re
        if re.match(r"^(\d+|[a-z]|[ivxl]+)\.", text, re.IGNORECASE):
            return True
        
        return False
    
    def _build_hierarchy(self, elements: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        Build parent-child hierarchy for elements
        Currently returns flat structure; can be enhanced for nested relationships
        """
        # For Phase 1, return flat structure
        # Phase 2 will enhance with actual hierarchy detection
        for element in elements:
            element["parent_id"] = None
            element["section_path"] = f"Page {element['page_number']}"
        
        return elements

    def _extract_docx_native(self, file_path: str) -> tuple[List[Dict[str, Any]], str]:
        """Fallback for DOCX files using python-docx"""
        try:
            doc = docx.Document(file_path)
            elements = []
            full_text_parts = []
            order_index = 0
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                full_text_parts.append(text)
                elements.append({
                    "element_id": f"native_text_{order_index}",
                    "element_type": "text",
                    "page_number": 1, 
                    "order_index": order_index,
                    "bounding_box": None,
                    "content": {"text": text},
                    "section_path": "Document Root"
                })
                order_index += 1
            
            # Extract tables too
            for table in doc.tables:
                table_data = {"headers": [], "rows": []}
                for i, row in enumerate(table.rows):
                    row_data = [{"text": cell.text.strip(), "row_span": 1, "col_span": 1} for cell in row.cells]
                    if i == 0:
                        table_data["headers"] = [[c.text.strip() for c in row.cells]]
                    table_data["rows"].append(row_data)
                
                elements.append({
                    "element_id": f"native_table_{order_index}",
                    "element_type": "table",
                    "page_number": 1,
                    "order_index": order_index,
                    "bounding_box": None,
                    "content": table_data,
                    "section_path": "Document Root"
                })
                order_index += 1
                full_text_parts.append(f"[Table with {len(table.rows)} rows]")

            return elements, "\n".join(full_text_parts)
        except Exception as e:
            logger.error(f"Native DOCX extraction failed: {e}")
            return [], ""
    
    async def _extract_pdf_fallback(self, file_path: str, reader: Optional[PdfReader] = None) -> tuple[list, str]:
        """
        Fallback PDF extraction using pypdf when Document AI fails
        
        Returns:
            tuple: (elements list, full_text)
        """
        try:
            logger.info(f"Attempting PyPDF fallback extraction for: {file_path}")
            
            if not reader:
                with open(file_path, 'rb') as f:
                    # Read into memory to avoid issues with closed files
                    content = f.read()
                    reader = PdfReader(io.BytesIO(content))
            
            elements = []
            full_text = ""
            
            for page_num, page in enumerate(reader.pages, 1):
                page_text = page.extract_text()
                
                if page_text.strip():
                    # Create element matching Document AI schema
                    elements.append({
                        "element_id": f"pypdf_page_{page_num}",
                        "element_type": "page",
                        "page_number": page_num,
                        "order_index": page_num - 1,  # 0-indexed
                        "content": {"text": page_text},  # Nested in content dict
                        "section_path": f"Page {page_num}",
                        "bounding_box": None  # PyPDF doesn't provide bounding boxes
                    })
                    full_text += page_text + "\n\n"
            
            logger.info(f"PyPDF extracted {len(elements)} pages, {len(full_text)} characters")
            return elements, full_text
                
        except Exception as e:
            logger.error(f"PyPDF fallback extraction failed: {e}", exc_info=True)
            return [], ""
