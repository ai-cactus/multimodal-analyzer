"""
Iterative Document Analysis Service - Automatic refinement workflow
"""
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, update
from datetime import datetime
from typing import Optional, Tuple, List, Dict, Any
from uuid import UUID, uuid4
import os
import json
import tempfile
from docx import Document as DocxDocument
from docx.shared import Inches

from app.models import UploadedDocument, Finding
from app.services.analysis_service import DocumentAnalysisService
from app.services.cost_tracker import CostTracker
from app.config import get_settings
from app.utils.logger import get_logger

logger = get_logger(__name__)


class IterativeAnalysisService:
    """
    Wraps DocumentAnalysisService to provide automatic iterative refinement
    """
    
    def __init__(self):
        self.analysis_service = DocumentAnalysisService()
        self.settings = get_settings()
    
    async def analyze_document_iterative(
        self,
        db: AsyncSession,
        document_id: UUID,
        file_path: str,
        compliance_body: str,
        compliance_section: Optional[str] = None,
        enable_iterations: bool = True
    ) -> UUID:
        """
        Execute iterative analysis workflow with automatic refinement
        
        Args:
            db: Database session
            document_id: Initial document ID
            file_path: Path to document file
            compliance_body: Compliance body name
            compliance_section: Optional section filter
            enable_iterations: Enable automatic re-analysis (default: True)
            
        Returns:
            Final document ID after all iterations
        """
        if not enable_iterations:
            # Run single-pass analysis
            await self.analysis_service.analyze_document(
                db, document_id, file_path, compliance_body, compliance_section
            )
            return document_id
        
        iteration = 0
        current_file_path = file_path
        current_doc_id = document_id
        parent_doc_id = document_id
        
        # Create top-level cost tracker to accumulate across iterations
        cumulative_cost_tracker = CostTracker(document_id=str(document_id))
        
        logger.info(f"Starting iterative analysis (max iterations: {self.settings.max_iterations})")
        
        while iteration < self.settings.max_iterations:
            logger.info(f"=== Iteration {iteration + 1}/{self.settings.max_iterations} ===")
            
            # Create per-iteration cost tracker
            iteration_cost_tracker = CostTracker(document_id=str(current_doc_id))
            
            # Run single analysis pass
            high_conf_count, medium_conf_count, redraft_content = await self._single_analysis_pass(
                db=db,
                document_id=current_doc_id,
                file_path=current_file_path,
                compliance_body=compliance_body,
                compliance_section=compliance_section,
                iteration=iteration,
                parent_doc_id=parent_doc_id,
                cost_tracker=iteration_cost_tracker,
            )
            
            # Merge iteration costs into cumulative tracker
            cumulative_cost_tracker.merge(iteration_cost_tracker)
            
            # Check convergence - both high AND medium confidence must be at/below thresholds
            converged = (
                high_conf_count <= self.settings.convergence_threshold
                and medium_conf_count <= self.settings.medium_convergence_threshold
            )
            
            logger.info(
                f"Iteration {iteration + 1} complete: "
                f"{high_conf_count} high-confidence findings, "
                f"{medium_conf_count} medium-confidence findings"
            )
            
            if converged:
                logger.info(f"✓ Converged at iteration {iteration + 1} - Document is compliant!")
                await self._update_iteration_history(
                    db, parent_doc_id, iteration, current_doc_id, 
                    high_conf_count, medium_conf_count, converged=True
                )
                break
            
            # Store iteration metadata
            await self._update_iteration_history(
                db, parent_doc_id, iteration, current_doc_id,
                high_conf_count, medium_conf_count, converged=False
            )
            
            # Check if we've hit max iterations
            if iteration >= self.settings.max_iterations - 1:
                logger.warning(
                    f"Max iterations ({self.settings.max_iterations}) reached without full convergence. "
                    f"Final document has {high_conf_count} high-confidence findings."
                )
                break
            
            # Prepare for next iteration
            if not redraft_content:
                logger.error("No redraft generated, cannot proceed to next iteration")
                break
            
            # Save redraft as file for next iteration
            try:
                current_file_path = await self._save_redraft_as_file(redraft_content, iteration + 1)
            except Exception as e:
                logger.error(f"Failed to save redraft for next iteration: {e}")
                break
            
            # Create new document record for next iteration
            current_doc_id = await self._create_iteration_document(
                db, parent_doc_id, iteration + 1, current_file_path
            )
            
            iteration += 1
        
        # Store cumulative cost tracking in parent document meta
        cost_summary = cumulative_cost_tracker.get_summary()
        result = await db.execute(
            select(UploadedDocument).where(UploadedDocument.id == parent_doc_id)
        )
        parent_doc = result.scalar_one()
        updated_meta = parent_doc.meta or {}
        updated_meta["cost_tracking"] = cost_summary
        await db.execute(
            update(UploadedDocument)
            .where(UploadedDocument.id == parent_doc_id)
            .values(meta=updated_meta)
        )
        await db.commit()
        
        logger.info(
            f"Iterative analysis complete. Total iterations: {iteration + 1}, "
            f"Final document ID: {current_doc_id}",
            extra={"cumulative_cost": cost_summary},
        )
        
        return current_doc_id
    
    async def _single_analysis_pass(
        self,
        db: AsyncSession,
        document_id: UUID,
        file_path: str,
        compliance_body: str,
        compliance_section: Optional[str],
        iteration: int,
        parent_doc_id: UUID,
        cost_tracker: Optional[CostTracker] = None,
    ) -> Tuple[int, int, Optional[str]]:
        """
        Execute a single analysis pass
        
        Returns:
            (high_confidence_count, medium_confidence_count, redraft_content)
        """
        # Run the standard analysis workflow
        await self.analysis_service.analyze_document(
            db, document_id, file_path, compliance_body, compliance_section,
            cost_tracker=cost_tracker,
        )
        
        # Extract results
        result = await db.execute(
            select(UploadedDocument).where(UploadedDocument.id == document_id)
        )
        doc = result.scalar_one()
        
        high_conf_findings = doc.meta.get('high_confidence_findings', [])
        medium_conf_findings = doc.meta.get('medium_confidence_findings', [])
        redraft_content = doc.meta.get('redraft_content', '')
        
        # Store findings in database with unique IDs
        await self._store_findings(
            db, document_id, iteration, high_conf_findings, medium_conf_findings
        )
        
        return len(high_conf_findings), len(medium_conf_findings), redraft_content
    
    async def _store_findings(
        self,
        db: AsyncSession,
        document_id: UUID,
        iteration: int,
        high_conf: List[Dict],
        medium_conf: List[Dict]
    ):
        """Store findings in database with unique error IDs"""
        
        all_findings = []
        
        # Process high confidence findings
        for finding in high_conf:
            all_findings.append({
                'finding': finding,
                'confidence': finding.get('confidence', 0.95),
                'is_high': True
            })
        
        # Process medium confidence findings
        for finding in medium_conf:
            all_findings.append({
                'finding': finding,
                'confidence': finding.get('confidence', 0.85),
                'is_high': False
            })
        
        # Create Finding records
        for idx, item in enumerate(all_findings):
            finding_data = item['finding']
            
            finding_record = Finding(
                finding_id=f"err-{uuid4().hex[:6]}",
                document_id=document_id,
                iteration_number=iteration,
                confidence_score=item['confidence'],
                status='active',
                model_name=finding_data.get('model_name', 'unknown'),
                error_type=finding_data.get('rule', finding_data.get('type', 'compliance_gap')),
                description=finding_data.get('description', ''),
                regulation_ref=finding_data.get('regulation_ref', ''),
                element_id=finding_data.get('element_id', ''),
                critique_trail=finding_data.get('critique_trail', []),
                meta=finding_data
            )
            db.add(finding_record)
        
        await db.commit()
        logger.info(f"Stored {len(all_findings)} findings for iteration {iteration}")
    
    async def _save_redraft_as_file(self, redraft_content: str, iteration: int) -> str:
        """
        Save redraft text as a DOCX file for next iteration
        
        Args:
            redraft_content: Text content of redraft
            iteration: Current iteration number
            
        Returns:
            Path to saved DOCX file
        """
        # Create temporary file
        temp_dir = tempfile.gettempdir()
        filename = f"redraft_iter_{iteration}_{uuid4().hex[:8]}.docx"
        file_path = os.path.join(temp_dir, filename)
        
        # Create DOCX document
        doc = DocxDocument()
        doc.add_heading(f'Redraft - Iteration {iteration}', 0)
        
        # Add redraft content (split by paragraphs)
        paragraphs = redraft_content.split('\n\n')
        for para in paragraphs:
            if para.strip():
                doc.add_paragraph(para.strip())
        
        doc.save(file_path)
        logger.info(f"Saved redraft to {file_path}")
        
        return file_path
    
    async def _create_iteration_document(
        self,
        db: AsyncSession,
        parent_doc_id: UUID,
        iteration: int,
        file_path: str
    ) -> UUID:
        """
        Create a new document record for the next iteration
        
        Args:
            db: Database session
            parent_doc_id: ID of the original parent document
            iteration: Next iteration number
            file_path: Path to redraft file
            
        Returns:
            New document ID
        """
        # Get parent document
        result = await db.execute(
            select(UploadedDocument).where(UploadedDocument.id == parent_doc_id)
        )
        parent_doc = result.scalar_one()
        
        # Create new document record
        new_doc = UploadedDocument(
            filename=f"{parent_doc.filename}_iter{iteration}.docx",
            compliance_body_id=parent_doc.compliance_body_id,
            compliance_section=parent_doc.compliance_section,
            status='processing',
            file_path=file_path,
            redraft_mode=parent_doc.redraft_mode,
            output_formats=parent_doc.output_formats,
            iteration_count=iteration,
            parent_document_id=parent_doc_id,
            iteration_history=[]
        )
        
        db.add(new_doc)
        await db.commit()
        await db.refresh(new_doc)
        
        logger.info(f"Created iteration {iteration} document: {new_doc.id}")
        
        return new_doc.id
    
    async def _update_iteration_history(
        self,
        db: AsyncSession,
        parent_doc_id: UUID,
        iteration: int,
        document_id: UUID,
        high_conf_count: int,
        medium_conf_count: int,
        converged: bool
    ):
        """Update the parent document's iteration history"""
        result = await db.execute(
            select(UploadedDocument).where(UploadedDocument.id == parent_doc_id)
        )
        parent_doc = result.scalar_one()
        
        history_entry = {
            'iteration': iteration,
            'document_id': str(document_id),
            'high_confidence_findings': high_conf_count,
            'medium_confidence_findings': medium_conf_count,
            'converged': converged,
            'timestamp': datetime.utcnow().isoformat()
        }
        
        iteration_history = parent_doc.iteration_history or []
        iteration_history.append(history_entry)
        
        await db.execute(
            update(UploadedDocument)
            .where(UploadedDocument.id == parent_doc_id)
            .values(iteration_history=iteration_history)
        )
        await db.commit()


# Singleton
_iterative_service = None


def get_iterative_analysis_service() -> IterativeAnalysisService:
    """Get or create singleton iterative analysis service"""
    global _iterative_service
    if _iterative_service is None:
        _iterative_service = IterativeAnalysisService()
    return _iterative_service
